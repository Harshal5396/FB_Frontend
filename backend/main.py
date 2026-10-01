import io
import os
import re
import shutil
import base64
from datetime import datetime
import pandas as pd
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ---------- DATABASE ----------
from database import SessionLocal
from models import WbRaw, Zv07Raw, MergedOutput


app = FastAPI(title="Excel Data Reconciliation API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
)


# ---------- FOLDER PATHS ----------
BASE_DATA_DIR = "data"
WB_UPLOADS_DIR = os.path.join(BASE_DATA_DIR, "wb_uploads")
RPA_OUTPUT_DIR = os.path.join(BASE_DATA_DIR, "rpa_output")
ARCHIVE_WB_DIR = os.path.join(BASE_DATA_DIR, "archive", "WB")
ARCHIVE_ZV07_DIR = os.path.join(BASE_DATA_DIR, "archive", "ZV07")
MERGED_OUTPUT_DIR = os.path.join(BASE_DATA_DIR, "merged_output")
LOGS_DIR = os.path.join(BASE_DATA_DIR, "logs")

for d in [WB_UPLOADS_DIR, RPA_OUTPUT_DIR, ARCHIVE_WB_DIR,
          ARCHIVE_ZV07_DIR, MERGED_OUTPUT_DIR, LOGS_DIR]:
    os.makedirs(d, exist_ok=True)


WB_REQUIRED = [
    ("Truck No", ["truckno", "trucknumber", "vehicleno"]),
    ("DateTime 2", ["datetime2", "datetime 2"]),
    ("Wt1", ["wt1"]),
    ("Wt2", ["wt2"]),
    ("Net Wt", ["netwt", "net wt", "netweight"]),
]

ZV_REQUIRED = [
    ("Bill Date", ["billdate"]),
    ("Plnt", ["plnt", "plant"]),
    ("Del No.", ["delno", "deliveryno"]),
    ("Distr Channel", ["distrchannel", "distributionchannel", "distchannel"]),
    ("Truck No.", ["truckno", "trucknumber", "vehicleno"]),
    ("Funct", ["funct"]),
]


# ---------- Helpers ----------

def _clean(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def _find_col(df, keywords):
    for col in df.columns:
        c = _clean(col)
        for kw in keywords:
            if kw in c:
                return col
    return None


def _resolve_required(df, required, label):
    resolved, missing = {}, []
    for logical, keywords in required:
        col = _find_col(df, keywords)
        if col is None:
            missing.append(logical)
        else:
            resolved[logical] = col
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Required column(s) not found in {label}: {missing}. Available: {list(df.columns)}",
        )
    return resolved


def _read_excel(file_bytes, filename):
    try:
        df_raw = pd.read_excel(io.BytesIO(file_bytes), header=None, engine="openpyxl")
        header_keywords = [
            "truck", "sno", "rstno", "billdate", "bill date",
            "supplier", "material", "grossweight", "netweight",
            "delno", "plnt", "plant", "distr", "funct", "wt1", "wt2",
            "datetime", "datetime2",
        ]
        header_idx = 0
        for idx, row in df_raw.iterrows():
            row_str = " ".join(str(v).lower() for v in row.values if pd.notna(v))
            if any(kw in row_str for kw in header_keywords):
                header_idx = idx
                break

        if header_idx > 0:
            df = pd.read_excel(io.BytesIO(file_bytes), skiprows=header_idx, engine="openpyxl")
        else:
            df = pd.read_excel(io.BytesIO(file_bytes), engine="openpyxl")

        df.columns = [str(c).strip() for c in df.columns]
        df = df.dropna(how="all").reset_index(drop=True)
        return df
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid Excel file '{filename}'. Error: {str(e)}")


def _extract_date(value):
    if pd.isna(value):
        return pd.NaT
    if isinstance(value, (pd.Timestamp, datetime)):
        return pd.Timestamp(value).normalize()
    s = str(value).strip()
    if not s:
        return pd.NaT
    formats = [
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
        "%m/%d/%Y %H:%M:%S", "%m/%d/%Y",
    ]
    for fmt in formats:
        try:
            return pd.Timestamp(datetime.strptime(s, fmt)).normalize()
        except ValueError:
            continue
    try:
        return pd.Timestamp(pd.to_datetime(s, dayfirst=True, errors="raise")).normalize()
    except Exception:
        return pd.NaT


def _normalize_truck(value):
    if pd.isna(value):
        return ""
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def _resolve_wb_duplicates(wb):
    drop_idx = []
    for _, g in wb.groupby(["_date", "_truck"]):
        n = len(g)
        if n < 2 or n >= 4:
            continue
        w1 = g["_wt1_num"]
        w2 = g["_wt2_num"]
        if w1.isna().all() or w2.isna().all():
            continue
        min_wt1_rows = set(g.index[w1 == w1.min()])
        max_wt2_rows = [i for i in g.index[w2 == w2.max()]]
        common = [i for i in max_wt2_rows if i in min_wt1_rows]
        if common:
            keep = common[0]
            drop_idx.extend([i for i in g.index if i != keep])
    if drop_idx:
        wb = wb.drop(index=drop_idx)
    return wb


def _get_report_date(wb_df, dt_col):
    try:
        dt_series = wb_df[dt_col].dropna()
        if len(dt_series) > 0:
            first = _extract_date(dt_series.iloc[0])
            if pd.notna(first):
                return first
    except Exception:
        pass
    return pd.Timestamp.today().normalize()


# ---------- DB Conversion Helpers ----------

def _ts_to_date(ts):
    if ts is None or pd.isna(ts):
        return None
    try:
        return ts.date()
    except Exception:
        return None


def _value_to_datetime(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    try:
        ts = pd.to_datetime(value, errors="coerce")
        if pd.notna(ts):
            return ts.to_pydatetime()
    except Exception:
        return None
    return None


def _safe_str(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    s = str(value).strip()
    return s if s else None


def _safe_float(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def _safe_int(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


# ---------- DB Save Functions ----------

def _save_wb_to_db(wb_df, wb_truck_col, wb_dt_col, wb_wt1_col, wb_wt2_col, wb_net_col):
    db = SessionLocal()
    try:
        records = []
        for _, row in wb_df.iterrows():
            records.append(WbRaw(
                date=_ts_to_date(_extract_date(row[wb_dt_col])),
                truck_no=_safe_str(row[wb_truck_col]),
                date_time_2=_value_to_datetime(row[wb_dt_col]),
                wt1=_safe_float(row[wb_wt1_col]),
                wt2=_safe_float(row[wb_wt2_col]),
                net_wt=_safe_float(row[wb_net_col]),
            ))
        db.add_all(records)
        db.commit()
        print(f"[DB] WB saved: {len(records)} rows")
    except Exception as e:
        db.rollback()
        print(f"[DB ERROR] WB save failed: {e}")
    finally:
        db.close()


def _save_zv07_to_db(zv, zv_bill_col, zv_plnt_col, zv_del_col,
                     zv_dist_col, zv_truck_col, zv_funct_col, file_date):
    db = SessionLocal()
    try:
        fd = _ts_to_date(file_date) if file_date is not None else None
        records = []
        for _, row in zv.iterrows():
            records.append(Zv07Raw(
                bill_date=_ts_to_date(_extract_date(row[zv_bill_col])),
                plant=_safe_int(row[zv_plnt_col]),
                del_no=_safe_str(row[zv_del_col]),
                distr_channel=_safe_str(row[zv_dist_col]),
                truck_no=_safe_str(row[zv_truck_col]),
                funct=_safe_str(row[zv_funct_col]),
                file_date=fd,
            ))
        db.add_all(records)
        db.commit()
        print(f"[DB] ZV07 saved: {len(records)} rows")
    except Exception as e:
        db.rollback()
        print(f"[DB ERROR] ZV07 save failed: {e}")
    finally:
        db.close()


def _save_merged_to_db(all_final):
    db = SessionLocal()
    try:
        records = []
        for _, row in all_final.iterrows():
            records.append(MergedOutput(
                bill_date=_ts_to_date(_extract_date(row.get("Bill Date"))),
                date_time_2=_value_to_datetime(row.get("DateTime 2")),
                plant=_safe_int(row.get("Plant")),
                distr_channel=_safe_str(row.get("Distr Channel")),
                funct=_safe_str(row.get("Funct")),
                status=_safe_str(row.get("Status")),
                truck_no=_safe_str(row.get("Truck No.")),
                wt2=_safe_float(row.get("Wt2")),
                net_wt=_safe_float(row.get("Net Wt")),
                del_no=_safe_str(row.get("Del No.")),
            ))
        db.add_all(records)
        db.commit()
        print(f"[DB] Merged saved: {len(records)} rows")
    except Exception as e:
        db.rollback()
        print(f"[DB ERROR] Merged save failed: {e}")
    finally:
        db.close()


# ---------- Excel Styling ----------

def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)


def _style_grouped_matched_sheet(ws, df, group_cols, highlight_cols):
    yellow_fill = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
    blue_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    dark = Side(border_style="medium", color="404040")
    bottom_border = Border(bottom=dark)
    header_row = [cell.value for cell in ws[1]]
    total_cols = ws.max_column
    styled_indexes = list(range(1, total_cols + 1))

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 40)

    if len(df) == 0:
        return

    truck_col_idx = header_row.index("Truck No.") + 1 if "Truck No." in header_row else None

    current_group_rows = []
    groups = []
    for r_idx in range(2, ws.max_row + 1):
        cell_value = ws.cell(row=r_idx, column=truck_col_idx).value if truck_col_idx else None
        if cell_value is not None and str(cell_value).strip() != "":
            if current_group_rows:
                groups.append(current_group_rows)
            current_group_rows = [r_idx]
        else:
            if current_group_rows:
                current_group_rows.append(r_idx)
    if current_group_rows:
        groups.append(current_group_rows)

    for idx, group in enumerate(groups):
        fill = yellow_fill if idx % 2 == 0 else blue_fill
        for r_idx in group:
            for col_idx in styled_indexes:
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.fill = fill
        last_row = group[-1]
        for col_idx in styled_indexes:
            cell = ws.cell(row=last_row, column=col_idx)
            cell.border = bottom_border


def _group_matched_df(df: pd.DataFrame, group_cols) -> pd.DataFrame:
    if df.empty:
        return df
    df = df.reset_index(drop=True).copy()
    df = df.drop_duplicates(subset=["Truck No.", "Wt2", "Del No."], keep="first").reset_index(drop=True)
    df["_group_id"] = df.groupby(group_cols, dropna=False).ngroup()
    df = df.sort_values("_group_id", kind="stable").reset_index(drop=True)
    first_row_of_group = df.groupby("_group_id").head(1).index.tolist()
    for col in group_cols:
        df[col] = df.apply(lambda row: row[col] if row.name in first_row_of_group else "", axis=1)
    df = df.drop(columns=["_group_id"])
    return df


def _df_to_excel_bytes(df, sheet_name, use_group_style=False, group_cols=None, highlight_cols=None):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]
        if use_group_style and group_cols and highlight_cols:
            _style_grouped_matched_sheet(ws, df, group_cols, highlight_cols)
        else:
            _style_header(ws)
    buf.seek(0)
    return buf.getvalue()


def _build_diagnostics(zv, wb, merged, zv_truck_col):
    wb_keys = set(zip(wb["_date"], wb["_truck"]))
    merged_keys = set(zip(merged["_date"], merged["_truck"]))
    unmatched = []
    zv_keys = zv.drop_duplicates(subset=["_date", "_truck"])
    for _, r in zv_keys.iterrows():
        key = (r["_date"], r["_truck"])
        if key in merged_keys:
            continue
        digits = re.sub(r"\D", "", r["_truck"])
        suggestions = []
        if len(digits) >= 4:
            same_day = wb[wb["_date"] == r["_date"]]["_truck"].unique()
            suggestions = sorted(t for t in same_day if re.sub(r"\D", "", t).endswith(digits[-4:]))[:5]
        unmatched.append({
            "truck": str(r[zv_truck_col]),
            "date": r["_date"].strftime("%d-%m-%Y"),
            "del_count": int((zv["_truck"] == r["_truck"]).sum()),
            "reason": "Truck No. not present in Weight Bridge for this date",
            "possible_wb_typo_matches": suggestions,
        })
    return {
        "zv_trucks_after_filter": int(zv["_truck"].nunique()),
        "trucks_in_merged_output": int(merged["_truck"].nunique()),
        "zv_trucks_not_in_wb": unmatched,
    }


# ---------- MERGE LOGIC (wrapper) ----------

def _perform_merge(wb_bytes, zv_bytes, wb_filename, zv_filename):
    """
    Runs the full merge logic (same as before) and returns dict with all outputs.
    This is the exact same business logic, just extracted so both
    /process-excel and /trigger-merge can use it.
    """
    wb_df = _read_excel(wb_bytes, wb_filename)
    zv_df = _read_excel(zv_bytes, zv_filename)

    wb_map = _resolve_required(wb_df, WB_REQUIRED, "Weight Bridge file")
    zv_map = _resolve_required(zv_df, ZV_REQUIRED, "ZV07 file")

    wb_truck_col = wb_map["Truck No"]
    wb_dt_col = wb_map["DateTime 2"]
    wb_wt1_col = wb_map["Wt1"]
    wb_wt_col = wb_map["Wt2"]
    wb_net_col = wb_map["Net Wt"]

    zv_bill_col = zv_map["Bill Date"]
    zv_plnt_col = zv_map["Plnt"]
    zv_del_col = zv_map["Del No."]
    zv_dist_col = zv_map["Distr Channel"]
    zv_truck_col = zv_map["Truck No."]
    zv_funct_col = zv_map["Funct"]

    _save_wb_to_db(wb_df, wb_truck_col, wb_dt_col, wb_wt1_col, wb_wt_col, wb_net_col)

    # ZV07 filter
    zv = zv_df.copy()
    zv["_plnt_num"] = pd.to_numeric(zv[zv_plnt_col], errors="coerce")
    zv = zv[zv["_plnt_num"].between(7100, 7400, inclusive="both")]
    zv = zv[zv[zv_dist_col].astype(str).str.upper().str.strip() == "DO"]
    zv = zv[zv[zv_funct_col].astype(str).str.upper().str.strip() == "ZZ"]

    if zv.empty:
        raise HTTPException(status_code=400, detail="No records in ZV07 after filters (Plant 7100-7400, DO, ZZ).")

    try:
        _fd_series = zv[zv_bill_col].dropna()
        file_date_val = _extract_date(_fd_series.iloc[0]) if len(_fd_series) > 0 else pd.Timestamp.today()
    except Exception:
        file_date_val = pd.Timestamp.today()
    _save_zv07_to_db(zv, zv_bill_col, zv_plnt_col, zv_del_col,
                     zv_dist_col, zv_truck_col, zv_funct_col, file_date_val)

    # WB prepare
    wb = wb_df.copy()
    wb["_date"] = wb[wb_dt_col].apply(_extract_date)
    wb["_truck"] = wb[wb_truck_col].apply(_normalize_truck)
    wb = wb[(wb["_date"].notna()) & (wb["_truck"] != "")].copy()

    wb["_wt1_num"] = pd.to_numeric(wb[wb_wt1_col], errors="coerce")
    wb["_wt2_num"] = pd.to_numeric(wb[wb_wt_col], errors="coerce")
    wb = _resolve_wb_duplicates(wb)

    wb["_wb_dup_count"] = wb.groupby(["_date", "_truck"])["_date"].transform("size")
    wb_dup_keys = (
        wb[wb["_wb_dup_count"] > 1][["_date", "_truck", "_wb_dup_count"]]
        .drop_duplicates(subset=["_date", "_truck"])
        .rename(columns={"_wb_dup_count": "WB_Count"})
    )

    # ZV07 prepare
    zv["_date"] = zv[zv_bill_col].apply(_extract_date)
    zv["_truck"] = zv[zv_truck_col].apply(_normalize_truck)
    zv = zv[(zv["_date"].notna()) & (zv["_truck"] != "")]

    if zv.empty:
        raise HTTPException(status_code=400, detail="No valid ZV07 records after parsing dates.")

    wb_cols = wb[["_date", "_truck", wb_dt_col, wb_wt_col, wb_net_col]].rename(
        columns={wb_dt_col: "DateTime 2", wb_wt_col: "Wt2", wb_net_col: "Net Wt"}
    )
    zv_cols = zv[["_date", "_truck", zv_bill_col, zv_plnt_col, zv_del_col,
                  zv_dist_col, zv_truck_col, zv_funct_col]].rename(columns={
        zv_bill_col: "Bill Date", zv_plnt_col: "Plnt", zv_del_col: "Del No.",
        zv_dist_col: "Distr Channel", zv_truck_col: "ZV_Truck", zv_funct_col: "Funct",
    })

    merged = pd.merge(zv_cols, wb_cols, on=["_date", "_truck"], how="inner")
    if merged.empty:
        raise HTTPException(status_code=400, detail="No matches found between files.")

    diagnostics = _build_diagnostics(zv, wb, merged, zv_truck_col)

    merged = merged.merge(wb_dup_keys, on=["_date", "_truck"], how="left")
    merged["_is_wb_dup"] = merged["WB_Count"].notna()
    merged["Status"] = merged["_is_wb_dup"].apply(lambda x: "Exception - Duplicate" if x else "Matched")

    matched_df = merged[~merged["_is_wb_dup"]].copy()
    exception_df = merged[merged["_is_wb_dup"]].copy()

    NOT_IN_WB_REASON = "Truck Number not found in WB"
    NOT_IN_WB_STATUS = f"Exception - {NOT_IN_WB_REASON}"

    wb_keys_df = wb[["_date", "_truck"]].drop_duplicates()
    not_found_df = zv_cols.merge(wb_keys_df, on=["_date", "_truck"], how="left", indicator=True)
    not_found_df = not_found_df[not_found_df["_merge"] == "left_only"].drop(columns=["_merge"])
    not_found_df = not_found_df.drop_duplicates(subset=["_date", "_truck", "Del No."]).copy()
    not_found_df["DateTime 2"] = ""
    not_found_df["Wt2"] = ""
    not_found_df["Net Wt"] = ""
    not_found_df["Status"] = NOT_IN_WB_STATUS

    final_cols = [
        "Bill Date", "DateTime 2", "Plnt", "Distr Channel", "Funct", "Status",
        "ZV_Truck", "Wt2", "Net Wt", "Del No."
    ]
    rename_map = {"Plnt": "Plant", "ZV_Truck": "Truck No."}

    matched_final = matched_df[final_cols].rename(columns=rename_map).reset_index(drop=True)
    exception_full = exception_df[final_cols].rename(columns=rename_map).reset_index(drop=True)
    not_found_full = not_found_df[final_cols].rename(columns=rename_map).reset_index(drop=True)
    not_found_full = not_found_full.sort_values(["Truck No.", "Del No."], kind="stable").reset_index(drop=True)
    exception_full = pd.concat([exception_full, not_found_full], ignore_index=True)

    group_cols = ["Truck No.", "Wt2", "Net Wt"]
    matched_grouped = _group_matched_df(matched_final, group_cols)

    matched_plain = matched_final.sort_values(
        ["Truck No.", "Wt2", "Del No."], kind="stable"
    ).reset_index(drop=True)
    all_final = pd.concat([matched_plain, exception_full], ignore_index=True)

    _save_merged_to_db(all_final)

    exception_simple = (
        exception_df[["_date", "ZV_Truck", "WB_Count"]]
        .drop_duplicates(subset=["_date", "ZV_Truck"])
        .rename(columns={"ZV_Truck": "Truck No."})
        .reset_index(drop=True)
    )
    exception_simple["Date"] = exception_simple["_date"].dt.strftime("%d-%m-%Y")
    exception_simple["Reason"] = "Exception - Duplicate"
    exception_simple["WB_Count"] = exception_simple["WB_Count"].astype(int)
    exception_simple = exception_simple[["Date", "Truck No.", "WB_Count", "Reason"]]

    not_found_simple = (
        not_found_df[["_date", "ZV_Truck"]]
        .drop_duplicates(subset=["_date", "ZV_Truck"])
        .rename(columns={"ZV_Truck": "Truck No."})
        .reset_index(drop=True)
    )
    not_found_simple["Date"] = not_found_simple["_date"].dt.strftime("%d-%m-%Y")
    not_found_simple["WB_Count"] = 0
    not_found_simple["Reason"] = NOT_IN_WB_REASON
    not_found_simple = not_found_simple[["Date", "Truck No.", "WB_Count", "Reason"]]

    exception_simple = pd.concat([exception_simple, not_found_simple], ignore_index=True)

    report_date = _get_report_date(wb_df, wb_dt_col)
    date_str = report_date.strftime("%d-%m-%Y")

    highlight_cols = ["Truck No.", "Wt2", "Net Wt", "Del No."]

    all_buf = io.BytesIO()
    with pd.ExcelWriter(all_buf, engine="openpyxl") as writer:
        all_final.to_excel(writer, sheet_name="All_Records", index=False)
        _style_header(writer.sheets["All_Records"])
        matched_grouped.to_excel(writer, sheet_name="Matched", index=False)
        _style_grouped_matched_sheet(
            writer.sheets["Matched"], matched_grouped,
            group_cols=group_cols, highlight_cols=highlight_cols
        )
        exception_simple.to_excel(writer, sheet_name="Exceptions", index=False)
        _style_header(writer.sheets["Exceptions"])
    all_buf.seek(0)
    all_bytes = all_buf.getvalue()

    matched_bytes = _df_to_excel_bytes(
        matched_grouped, "Matched",
        use_group_style=True, group_cols=group_cols, highlight_cols=highlight_cols
    )
    exception_bytes = _df_to_excel_bytes(exception_simple, "Exceptions")

    return {
        "all_bytes": all_bytes,
        "matched_bytes": matched_bytes,
        "exception_bytes": exception_bytes,
        "all_count": int(len(all_final)),
        "matched_count": int(len(matched_final)),
        "exception_count": int(len(exception_simple)),
        "date_str": date_str,
        "diagnostics": diagnostics,
    }


# ---------- Folder Helpers ----------

def _month_folder(ddmmyyyy: str) -> str:
    """'15-09-2026' -> '2026-09'"""
    try:
        dt = datetime.strptime(ddmmyyyy, "%d-%m-%Y")
        return dt.strftime("%Y-%m")
    except Exception:
        return datetime.today().strftime("%Y-%m")


# ---------- Routes ----------

@app.get("/")
def read_root():
    return {"status": "running", "message": "Reconciliation API is active"}


# ----- OLD endpoint (manual testing — keep same) -----
@app.post("/process-excel")
async def process_excel(
    wb_file: UploadFile = File(...),
    zv07_file: UploadFile = File(...),
):
    wb_bytes = await wb_file.read()
    zv_bytes = await zv07_file.read()

    if not wb_bytes:
        raise HTTPException(status_code=400, detail="Weight Bridge file is empty.")
    if not zv_bytes:
        raise HTTPException(status_code=400, detail="ZV07 file is empty.")

    result = _perform_merge(wb_bytes, zv_bytes, wb_file.filename, zv07_file.filename)
    date_str = result["date_str"]

    return JSONResponse(content={
        "all_file": {
            "name": f"AllRecords_{date_str}.xlsx",
            "content_b64": base64.b64encode(result["all_bytes"]).decode("utf-8"),
            "count": result["all_count"],
        },
        "matched_file": {
            "name": f"Reconciliation_{date_str}.xlsx",
            "content_b64": base64.b64encode(result["matched_bytes"]).decode("utf-8"),
            "count": result["matched_count"],
        },
        "exception_file": {
            "name": f"Exceptions_{date_str}.xlsx",
            "content_b64": base64.b64encode(result["exception_bytes"]).decode("utf-8"),
            "count": result["exception_count"],
        },
        "diagnostics": result["diagnostics"],
    })


# ----- NEW: User WB upload -----
@app.post("/upload-wb")
async def upload_wb(wb_file: UploadFile = File(...)):
    """User uploads WB file. Saved to data/wb_uploads/ with date-based name."""
    wb_bytes = await wb_file.read()
    if not wb_bytes:
        raise HTTPException(status_code=400, detail="WB file is empty.")

    # Peek the file to find the date
    try:
        peek_df = _read_excel(wb_bytes, wb_file.filename)
        wb_map = _resolve_required(peek_df, WB_REQUIRED, "Weight Bridge file")
        date_series = peek_df[wb_map["DateTime 2"]].dropna()
        if len(date_series) > 0:
            report_date = _extract_date(date_series.iloc[0])
        else:
            report_date = pd.Timestamp.today().normalize()
    except HTTPException:
        raise
    except Exception:
        report_date = pd.Timestamp.today().normalize()

    date_str = report_date.strftime("%d-%m-%Y")
    save_name = f"WB_{date_str}.xlsx"
    save_path = os.path.join(WB_UPLOADS_DIR, save_name)

    with open(save_path, "wb") as f:
        f.write(wb_bytes)

    print(f"[UPLOAD] WB saved: {save_path}")

    return {
        "status": "success",
        "message": f"WB file uploaded for {date_str}. Merged file will be ready by morning.",
        "date": date_str,
        "saved_as": save_name,
    }


# ----- NEW: RPA calls this after ZV07 download -----
@app.post("/trigger-merge")
async def trigger_merge(payload: dict):
    """
    RPA calls this after downloading ZV07 from SAP.
    Body: {"date": "15-09-2026"}
    """
    date_str = payload.get("date")
    if not date_str:
        # Fallback to today's date
        date_str = datetime.today().strftime("%d-%m-%Y")

    wb_path = os.path.join(WB_UPLOADS_DIR, f"WB_{date_str}.xlsx")
    zv_path = os.path.join(RPA_OUTPUT_DIR, f"ZV07_{date_str}.xlsx")

    if not os.path.exists(wb_path):
        raise HTTPException(status_code=404, detail=f"WB file not found for {date_str}")
    if not os.path.exists(zv_path):
        raise HTTPException(status_code=404, detail=f"ZV07 file not found for {date_str}")

    with open(wb_path, "rb") as f:
        wb_bytes = f.read()
    with open(zv_path, "rb") as f:
        zv_bytes = f.read()

    # Run merge (same logic)
    result = _perform_merge(wb_bytes, zv_bytes, f"WB_{date_str}.xlsx", f"ZV07_{date_str}.xlsx")

    # Save merged output (month-wise)
    month = _month_folder(date_str)
    merged_dir = os.path.join(MERGED_OUTPUT_DIR, month)
    os.makedirs(merged_dir, exist_ok=True)

    merged_path = os.path.join(merged_dir, f"Merged_{date_str}.xlsx")

    # Save all 3 sheets combined into one file
    with pd.ExcelWriter(merged_path, engine="openpyxl") as writer:
        # Re-read sheets from in-memory bytes
        all_df = pd.read_excel(io.BytesIO(result["all_bytes"]), sheet_name="All_Records")
        matched_df = pd.read_excel(io.BytesIO(result["matched_bytes"]), sheet_name="Matched")
        exc_df = pd.read_excel(io.BytesIO(result["exception_bytes"]), sheet_name="Exceptions")

        all_df.to_excel(writer, sheet_name="All_Records", index=False)
        _style_header(writer.sheets["All_Records"])
        matched_df.to_excel(writer, sheet_name="Matched", index=False)
        _style_grouped_matched_sheet(
            writer.sheets["Matched"], matched_df,
            group_cols=["Truck No.", "Wt2", "Net Wt"],
            highlight_cols=["Truck No.", "Wt2", "Net Wt", "Del No."]
        )
        exc_df.to_excel(writer, sheet_name="Exceptions", index=False)
        _style_header(writer.sheets["Exceptions"])

    print(f"[MERGE] Saved: {merged_path}")

    # Move WB to archive
    wb_archive_dir = os.path.join(ARCHIVE_WB_DIR, month)
    os.makedirs(wb_archive_dir, exist_ok=True)
    wb_archive_path = os.path.join(wb_archive_dir, f"WB_{date_str}.xlsx")
    shutil.move(wb_path, wb_archive_path)
    print(f"[ARCHIVE] WB moved: {wb_archive_path}")

    # Move ZV07 to archive
    zv_archive_dir = os.path.join(ARCHIVE_ZV07_DIR, month)
    os.makedirs(zv_archive_dir, exist_ok=True)
    zv_archive_path = os.path.join(zv_archive_dir, f"ZV07_{date_str}.xlsx")
    shutil.move(zv_path, zv_archive_path)
    print(f"[ARCHIVE] ZV07 moved: {zv_archive_path}")

    return {
        "status": "success",
        "message": f"Merge completed for {date_str}",
        "date": date_str,
        "merged_file": merged_path,
        "all_count": result["all_count"],
        "matched_count": result["matched_count"],
        "exception_count": result["exception_count"],
    }


# ----- NEW: Frontend status check -----
@app.get("/status/{date_str}")
async def get_status(date_str: str):
    """
    date_str format: '15-09-2026'
    Returns status of WB, ZV07, Merged files.
    """
    wb_path = os.path.join(WB_UPLOADS_DIR, f"WB_{date_str}.xlsx")
    month = _month_folder(date_str)
    zv_path = os.path.join(RPA_OUTPUT_DIR, f"ZV07_{date_str}.xlsx")
    merged_path = os.path.join(MERGED_OUTPUT_DIR, month, f"Merged_{date_str}.xlsx")

    # Also check archive (if merge already done)
    wb_archived = os.path.join(ARCHIVE_WB_DIR, month, f"WB_{date_str}.xlsx")
    zv_archived = os.path.join(ARCHIVE_ZV07_DIR, month, f"ZV07_{date_str}.xlsx")

    return {
        "date": date_str,
        "wb_uploaded": os.path.exists(wb_path) or os.path.exists(wb_archived),
        "zv07_downloaded": os.path.exists(zv_path) or os.path.exists(zv_archived),
        "merged_ready": os.path.exists(merged_path),
    }


# ----- NEW: Frontend download merged file -----
@app.get("/download/{date_str}")
async def download_merged(date_str: str):
    month = _month_folder(date_str)
    merged_path = os.path.join(MERGED_OUTPUT_DIR, month, f"Merged_{date_str}.xlsx")

    if not os.path.exists(merged_path):
        raise HTTPException(status_code=404, detail=f"Merged file not ready for {date_str}")

    return FileResponse(
        path=merged_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Merged_{date_str}.xlsx",
    )